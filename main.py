"""Orchestrator: scrape -> analyze -> Excel report.

Usage:
    python main.py                 # run all profiles in profiles.txt
    python main.py user1 user2     # run only these usernames
"""
import sys
import traceback

import anthropic
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import config
from scraper import scrape_profile
from analyzer import detect_brand_deals, detect_alcohol


def load_profiles() -> list[str]:
    lines = config.PROFILES_FILE.read_text().splitlines()
    return [l.strip().lstrip("@") for l in lines
            if l.strip() and not l.strip().startswith("#")]


def analyze_profile(username: str, client: anthropic.Anthropic) -> dict:
    posts = scrape_profile(username)

    deal_posts, deal_brands = [], set()
    alcohol_posts = []

    for post in posts:
        deal = detect_brand_deals(post)
        if deal["is_deal"]:
            deal_posts.append(post)
            deal_brands.update(deal["brands"])

        # Only videos can be scanned for alcohol.
        if post.get("video_url"):
            alc = detect_alcohol(post, client)
            if alc.get("alcohol"):
                alcohol_posts.append({**post, "_alc": alc})

    return {
        "username": username,
        "posts_analyzed": len(posts),
        "deal_count": len(deal_posts),
        "brands": sorted(deal_brands),
        "alcohol_count": len(alcohol_posts),
        "alcohol_urls": [p["url"] for p in alcohol_posts],
        "deal_urls": [p["url"] for p in deal_posts],
    }


def write_report(rows: list[dict]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    headers = ["Influencer", "Posts analyzed (12mo)", "Brand-deal posts",
               "Brands mentioned", "Alcohol videos", "Alcohol post URLs",
               "Brand-deal post URLs"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="222E3C")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for r in rows:
        ws.append([
            r["username"],
            r["posts_analyzed"],
            r["deal_count"],
            ", ".join(r["brands"]),
            r["alcohol_count"],
            "\n".join(r["alcohol_urls"]),
            "\n".join(r["deal_urls"]),
        ])

    widths = [22, 16, 14, 40, 12, 45, 45]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(config.OUTPUT_XLSX)
    print(f"\nReport written to {config.OUTPUT_XLSX}")


def main():
    if not config.APIFY_TOKEN or not config.ANTHROPIC_API_KEY:
        sys.exit("Set APIFY_TOKEN and ANTHROPIC_API_KEY in your .env first.")

    profiles = sys.argv[1:] or load_profiles()
    client = anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY)

    rows = []
    for i, username in enumerate(profiles, 1):
        print(f"[{i}/{len(profiles)}] {username} ...", flush=True)
        try:
            rows.append(analyze_profile(username, client))
        except Exception:  # noqa: BLE001 - keep going if one profile fails
            print(f"  FAILED: {username}")
            traceback.print_exc()
            rows.append({"username": username, "posts_analyzed": 0,
                         "deal_count": 0, "brands": [], "alcohol_count": 0,
                         "alcohol_urls": [], "deal_urls": []})

    write_report(rows)


if __name__ == "__main__":
    main()
