"""Export the HTML report (config.OUTPUT_HTML) to CSV/Excel.

Parses the generated HTML report (no need to re-scrape or re-run the
vision model) and writes, next to it:
  - <report>.csv   summary table, one row per profile
  - <report>.xlsx  Summary + Flagged Videos + Deals to Review + Brand Deals sheets

Category columns are read from the report's own header, so reports made
with any category set (and older reports) export correctly.

Usage:
    python export_report.py            # REPORT_NAME=... to pick another report
"""
import csv
import html
import re
import sys

import openpyxl
from openpyxl.styles import Font

import config

# Non-category columns of the summary table; every other column is a category.
FIXED_COLS = ["Posts", "Deals", "Brands", "Deals to review", "Flagged"]

SUMMARY_HEAD_RE = re.compile(r'<table class="summary">\s*<thead><tr>(.*?)</tr></thead>', re.S)
TH_RE = re.compile(r'<th>([^<]*)</th>')
SUMMARY_ROW_RE = re.compile(
    r'<tr><td class="uname"><a href="#[^"]*">@([^<]+)</a></td>'
    r'<td class="status status-(\w+)">[^<]*</td>'
    r'((?:<td[^>]*>\d+</td>)+)</tr>'
)
NUM_CELL_RE = re.compile(r'>(\d+)</td>')

CARD_RE = re.compile(
    r'<a name="([^"]*)"></a>\s*<section class="card">(.*?)</section>', re.S
)

REJECT_REASON_RE = re.compile(
    r'<p class="status-banner rejected">Auto-rejected — (.*?)</p>', re.S
)

SCRAPE_ERROR_RE = re.compile(
    r'<p class="status-banner error">Scrape failed — (.*?) \(check the username\)</p>', re.S
)

REVIEW_ROW_RE = re.compile(
    r'<tr class="review"><td>(.*?)</td><td>(.*?)</td><td><a href="([^"]*)"', re.S
)

VISION_GAP_RE = re.compile(
    r'<p class="status-banner error">⚠ (\d+) of (\d+) videos could not be', re.S
)

FLAG_ROW_RE = re.compile(
    r'<tr><td>(.*?)</td><td class="ev">(.*?)</td>'
    r'<td><a href="([^"]*)"[^>]*>view&nbsp;post ↗</a></td></tr>', re.S
)
BADGE_RE = re.compile(r'<span class="badge"[^>]*>([^<]*)</span>')

DEAL_CHIP_RE = re.compile(r'<span class="chip">([^<]*)</span>')
DEAL_LINK_RE = re.compile(r'<a href="([^"]*)"[^>]*>deal \d+ ↗</a>')


def unesc(s: str) -> str:
    # The source data was double-escaped upstream (e.g. "&amp;#x27;"),
    # so unescape repeatedly until it stabilizes.
    prev = None
    while prev != s:
        prev, s = s, html.unescape(s)
    return s.strip()


def parse_report(text: str):
    head = SUMMARY_HEAD_RE.search(text)
    cols = TH_RE.findall(head.group(1))[2:] if head else []  # after Influencer, Status
    cats = [c for c in cols if c not in FIXED_COLS]
    summary = []
    for m in SUMMARY_ROW_RE.finditer(text):
        vals = dict(zip(cols, map(int, NUM_CELL_RE.findall(m.group(3)))))
        summary.append({
            "username": unesc(m.group(1)),
            "status": m.group(2).capitalize(),
            "posts": vals.get("Posts", 0),
            "deals": vals.get("Deals", 0),
            "brands": vals.get("Brands", 0),
            "review_deals": vals.get("Deals to review", 0),
            "flagged": vals.get("Flagged", 0),
            "cats": {c: vals.get(c, 0) for c in cats},
        })

    reasons = {}
    gaps = {}
    flagged_videos, brand_deals, review_deals = [], [], []
    for card_m in CARD_RE.finditer(text):
        username, body = unesc(card_m.group(1)), card_m.group(2)

        reason_m = REJECT_REASON_RE.search(body)
        if reason_m:
            reasons[username] = unesc(reason_m.group(1))
        scrape_m = SCRAPE_ERROR_RE.search(body)
        if scrape_m:
            reasons[username] = f"scrape failed: {unesc(scrape_m.group(1))}"

        for term, brands, url in REVIEW_ROW_RE.findall(body):
            review_deals.append({"username": username, "matched": unesc(term),
                                 "brands": unesc(brands), "url": url})

        gap_m = VISION_GAP_RE.search(body)
        if gap_m:
            gaps[username] = f"{gap_m.group(1)}/{gap_m.group(2)} videos not analyzed"

        for row_m in FLAG_ROW_RE.finditer(body):
            badges_html, evidence, url = row_m.groups()
            flags = ", ".join(BADGE_RE.findall(badges_html))
            flagged_videos.append({
                "username": username,
                "flags": flags,
                "evidence": unesc(evidence),
                "url": url,
            })

        chips = ", ".join(unesc(c) for c in DEAL_CHIP_RE.findall(body))
        links = DEAL_LINK_RE.findall(body)
        for i, url in enumerate(links):
            brand_deals.append({
                "username": username,
                "deal_index": i + 1,
                "brands": chips if i == 0 else "",
                "url": url,
            })

    for r in summary:
        r["reason"] = reasons.get(r["username"], "")
        r["gaps"] = gaps.get(r["username"], "")

    return summary, flagged_videos, brand_deals, review_deals


def _cats(summary) -> list[str]:
    return list(summary[0]["cats"]) if summary else []


def write_csv(summary, path):
    cats = _cats(summary)
    header = ["Influencer", "Status", "Reason", "Analysis gaps", "Posts analyzed",
              "Brand deals", "Brands mentioned", "Deals to review"] + cats
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in summary:
            w.writerow([r["username"], r["status"], r["reason"], r["gaps"], r["posts"],
                        r["deals"], r["brands"], r["review_deals"]]
                       + [r["cats"][c] for c in cats])
    print(f"Wrote {path}")


def write_xlsx(summary, flagged_videos, brand_deals, review_deals, path):
    wb = openpyxl.Workbook()
    cats = _cats(summary)

    ws = wb.active
    ws.title = "Summary"
    header = ["Influencer", "Status", "Reason", "Analysis gaps", "Posts analyzed",
              "Brand deals", "Brands mentioned", "Deals to review", "Flagged videos"] + cats
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in summary:
        ws.append([r["username"], r["status"], r["reason"], r["gaps"], r["posts"], r["deals"],
                   r["brands"], r["review_deals"], r["flagged"]] + [r["cats"][c] for c in cats])
    ws.freeze_panes = "A2"
    widths = [22, 12, 34, 24, 18, 12, 16, 14, 14] + [12] * len(cats)
    for col, width in zip(ws.columns, widths):
        ws.column_dimensions[col[0].column_letter].width = width

    ws2 = wb.create_sheet("Flagged Videos")
    ws2.append(["Influencer", "Flags", "Evidence", "Post URL"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for r in flagged_videos:
        ws2.append([r["username"], r["flags"], r["evidence"], r["url"]])
    for col, width in zip(ws2.columns, [22, 20, 60, 40]):
        ws2.column_dimensions[col[0].column_letter].width = width

    ws_r = wb.create_sheet("Deals to Review")
    ws_r.append(["Influencer", "Matched", "Brands mentioned", "Post URL"])
    for cell in ws_r[1]:
        cell.font = Font(bold=True)
    for r in review_deals:
        ws_r.append([r["username"], r["matched"], r["brands"], r["url"]])
    for col, width in zip(ws_r.columns, [22, 20, 50, 40]):
        ws_r.column_dimensions[col[0].column_letter].width = width

    ws3 = wb.create_sheet("Brand Deals")
    ws3.append(["Influencer", "Deal #", "Brands mentioned", "Post URL"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for r in brand_deals:
        ws3.append([r["username"], r["deal_index"], r["brands"], r["url"]])
    for col, width in zip(ws3.columns, [22, 8, 50, 40]):
        ws3.column_dimensions[col[0].column_letter].width = width

    wb.save(path)
    print(f"Wrote {path}")


def main():
    usernames = {u.lstrip("@") for u in sys.argv[1:]}

    text = config.OUTPUT_HTML.read_text(encoding="utf-8")
    summary, flagged_videos, brand_deals, review_deals = parse_report(text)
    if not summary:
        raise SystemExit(f"No summary rows found — is {config.OUTPUT_HTML.name} up to date?")

    if usernames:
        missing = usernames - {r["username"] for r in summary}
        if missing:
            raise SystemExit(f"Not found in report: {', '.join(sorted(missing))}")
        summary = [r for r in summary if r["username"] in usernames]
        flagged_videos = [r for r in flagged_videos if r["username"] in usernames]
        brand_deals = [r for r in brand_deals if r["username"] in usernames]
        review_deals = [r for r in review_deals if r["username"] in usernames]
        stem = config.OUTPUT_HTML.with_name(config.OUTPUT_HTML.stem + "_selected")
    else:
        stem = config.OUTPUT_HTML

    write_csv(summary, stem.with_suffix(".csv"))
    write_xlsx(summary, flagged_videos, brand_deals, review_deals, stem.with_suffix(".xlsx"))


if __name__ == "__main__":
    main()
