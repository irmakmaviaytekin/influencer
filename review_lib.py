"""Data layer for the review GUI (no Streamlit in here, so it can be tested alone).

Reads finished runs (runs/<timestamp>.json), re-applies a sensitivity to the saved
observations, logs reviewer decisions, and exports a report for legal.
"""
from __future__ import annotations

import io
import json
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).parent
RUNS_DIR = ROOT / "runs"
DECISIONS_FILE = ROOT / "decisions" / "decisions.jsonl"

LEVELS = {"incidental": 1, "notable": 2, "heavy": 3}
LEVEL_NAMES = {v: k for k, v in LEVELS.items()}
DECISIONS = ["Approve", "Reject", "Needs more review"]


# --------------------------------------------------------------------------- runs
def list_runs() -> list[dict]:
    """Newest first. Only real run files (not partial logs)."""
    out = []
    for p in sorted(RUNS_DIR.glob("2*.json"), reverse=True):
        try:
            d = json.loads(p.read_text())
        except (OSError, json.JSONDecodeError):
            continue
        job = d.get("job") or {}
        out.append({
            "id": p.stem, "path": p,
            "started": d.get("started", ""),
            "job": job.get("name") or d.get("settings", {}).get("category_set", ""),
            "brand": job.get("brand"),
            "profiles": len(d.get("profiles", [])),
            "posts": d.get("posts", 0), "failed": d.get("posts_failed", 0),
            "model": d.get("settings", {}).get("vision_model", ""),
            "cost": d.get("vision_cost_usd", 0),
            "has_levels": any("levels" in x for x in d.get("per_post", [])),
            # Runs from before creator names were saved can't be reviewed per creator.
            "reviewable": any("profile" in x for x in d.get("per_post", [])),
        })
    return out


def load_run(run_id: str) -> dict:
    run = json.loads((RUNS_DIR / f"{run_id}.json").read_text())
    run["id"] = run_id
    return run


def categories_of(run: dict) -> dict:
    """{key: {label, description, min_level (name)}} for the run's criteria."""
    cats = (run.get("job") or {}).get("categories")
    if cats:
        return cats
    keys = []  # older runs: infer from what was recorded
    for p in run.get("per_post", []):
        for k in list(p.get("levels") or {}) + list(p.get("hits") or []):
            if k not in keys:
                keys.append(k)
    return {k: {"label": k.replace("_", " ").title(), "description": "", "min_level": "heavy"}
            for k in keys}


def default_mins(run: dict) -> dict[str, str]:
    return {k: c.get("min_level", "heavy") for k, c in categories_of(run).items()}


# ------------------------------------------------------------------ flags / sensitivity
def flags_for(run: dict, mins: dict[str, str]) -> list[dict]:
    """Every observation at or above `mins` (category -> level name), re-applied from the
    saved levels. Runs that predate saved levels fall back to their recorded yes/no flags."""
    out = []
    for p in run.get("per_post", []):
        if p.get("error"):
            continue
        prof = p.get("profile", "?")
        if "levels" in p:
            for cat, lvl in (p.get("levels") or {}).items():
                if lvl >= LEVELS.get(mins.get(cat, "heavy"), 3):
                    out.append({"profile": prof, "url": p["url"], "category": cat, "level": lvl,
                                "reason": (p.get("observations") or {}).get(cat, ""),
                                "thumbs": p.get("thumbs") or []})
        else:
            for cat in p.get("hits") or []:
                out.append({"profile": prof, "url": p["url"], "category": cat, "level": None,
                            "reason": (p.get("evidence") or {}).get(cat, ""),
                            "thumbs": p.get("thumbs") or []})
    return out


def profile_table(run: dict, flags: list[dict]) -> list[dict]:
    """One row per creator: how much was analyzed and what was flagged."""
    analyzed, failed = {}, {}
    for p in run.get("per_post", []):
        prof = p.get("profile", "?")
        if p.get("error"):
            failed[prof] = failed.get(prof, 0) + 1
        else:
            analyzed[prof] = analyzed.get(prof, 0) + 1
    by: dict[str, dict[str, int]] = {}
    posts: dict[str, set] = {}
    for f in flags:
        by.setdefault(f["profile"], {})
        by[f["profile"]][f["category"]] = by[f["profile"]].get(f["category"], 0) + 1
        posts.setdefault(f["profile"], set()).add(f["url"])
    names = list(dict.fromkeys(run.get("profiles", []) + list(analyzed) + list(failed)))
    rows = []
    for n in names:
        a, fl = analyzed.get(n, 0), failed.get(n, 0)
        rows.append({"creator": n, "analyzed": a, "not_analyzed": fl,
                     "flagged_posts": len(posts.get(n, ())), "by_category": by.get(n, {}),
                     "incomplete": fl > 0 and a + fl > 0 and fl / (a + fl) >= 0.1})
    return rows


# ----------------------------------------------------------------------- decisions
def log_decision(entry: dict) -> dict:
    """Append one reviewer decision. Append-only: the log is the audit trail."""
    entry = {"ts": datetime.now().isoformat(timespec="seconds"), **entry}
    DECISIONS_FILE.parent.mkdir(exist_ok=True)
    with DECISIONS_FILE.open("a", encoding="utf-8") as f:
        f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    return entry


def load_decisions(run_id: str | None = None) -> list[dict]:
    if not DECISIONS_FILE.exists():
        return []
    out = []
    for line in DECISIONS_FILE.read_text(encoding="utf-8").splitlines():
        if line.strip():
            try:
                d = json.loads(line)
            except json.JSONDecodeError:
                continue
            if run_id is None or d.get("run") == run_id:
                out.append(d)
    return out


def latest_by_creator(run_id: str) -> dict[str, dict]:
    """The most recent decision per creator for a run (later entries override earlier)."""
    return {d["creator"]: d for d in load_decisions(run_id)}


# --------------------------------------------------------------------------- export
def export_xlsx(run: dict, mins: dict[str, str]) -> bytes:
    """A report for legal/brand: decisions, every flag with its evidence frame and
    link, and the exact criteria used - so sign-off doesn't need the videos re-watched."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill

    flags = flags_for(run, mins)
    table = profile_table(run, flags)
    decisions = latest_by_creator(run["id"])
    dismissed = {(d["creator"], u) for d in decisions.values() for u in d.get("dismissed", [])}
    cats = categories_of(run)
    job = run.get("job") or {}
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E7EFEC")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Decisions"
    ws.append(["Creator", "Decision", "Reviewer", "Decided", "Note", "Flagged posts (open)",
               "Flags by category", "Posts not analyzed"])
    for r in table:
        d = decisions.get(r["creator"], {})
        open_flags = [f for f in flags if f["profile"] == r["creator"]
                      and (r["creator"], f["url"]) not in dismissed]
        ws.append([r["creator"], d.get("decision", "— not reviewed —"), d.get("reviewer", ""),
                   d.get("ts", ""), d.get("note", ""), len({f["url"] for f in open_flags}),
                   ", ".join(f"{cats.get(k, {}).get('label', k)} {n}" for k, n in r["by_category"].items()),
                   r["not_analyzed"] or ""])
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
    for col, w in zip("ABCDEFGH", (24, 20, 16, 20, 50, 18, 36, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    fs = wb.create_sheet("Flags")
    fs.append(["Creator", "Category", "Level", "Reason (model)", "Post link", "Reviewer dismissed", "Evidence frame"])
    for c in fs[1]:
        c.font, c.fill = bold, head_fill
    for col, w in zip("ABCDEFG", (22, 18, 12, 70, 44, 18, 26)):
        fs.column_dimensions[col].width = w
    row = 2
    for f in sorted(flags, key=lambda x: (x["profile"], -(x["level"] or 0))):
        fs.append([f["profile"], cats.get(f["category"], {}).get("label", f["category"]),
                   LEVEL_NAMES.get(f["level"], "(flagged)"), f["reason"], f["url"],
                   "yes" if (f["profile"], f["url"]) in dismissed else ""])
        fs.cell(row, 4).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (1, 2, 3, 5, 6):
            fs.cell(row, col).alignment = Alignment(vertical="top")
        thumb = next((ROOT / t for t in f["thumbs"] if (ROOT / t).exists()), None)
        if thumb:
            try:
                from PIL import Image as PILImage
                im = PILImage.open(thumb)
                im.thumbnail((170, 170))
                buf = io.BytesIO()
                im.save(buf, "JPEG")
                buf.seek(0)
                img = XLImage(buf)
                fs.add_image(img, f"G{row}")
                fs.row_dimensions[row].height = max(im.height * 0.75 + 4, 30)
            except Exception:  # noqa: BLE001 - a bad image must not break the export
                pass
        row += 1
    fs.freeze_panes = "A2"

    cs = wb.create_sheet("Criteria used")
    cs.append(["Setting", "Value"])
    for c in cs[1]:
        c.font, c.fill = bold, head_fill
    for k, v in [("Run", run["id"]), ("Job", job.get("name")), ("Brand", job.get("brand")),
                 ("Campaign", job.get("campaign")), ("Model", run.get("settings", {}).get("vision_model")),
                 ("Lookback (months)", job.get("months_lookback")),
                 ("Reels/videos only", job.get("videos_only")),
                 ("Posts analyzed / not analyzed",
                  f"{run.get('posts', 0) - run.get('posts_failed', 0)} / {run.get('posts_failed', 0)}")]:
        cs.append([k, "" if v is None else str(v)])
    cs.append([])
    cs.append(["Category", "Flagged at", "What counts"])
    for c in cs[cs.max_row]:
        c.font, c.fill = bold, head_fill
    for k, c in cats.items():
        cs.append([c.get("label", k), mins.get(k, c.get("min_level", "heavy")) + " or above",
                   c.get("description", "")])
    cs.column_dimensions["A"].width = 30
    cs.column_dimensions["B"].width = 24
    cs.column_dimensions["C"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
